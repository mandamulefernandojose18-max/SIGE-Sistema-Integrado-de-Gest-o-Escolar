import pytest
from apps.documentos.services.excel_service import gerar_caderneta_xlsx, gerar_pauta_xlsx
from apps.documentos.services.pdf_service import gerar_caderneta_pdf, gerar_pauta_pdf
import openpyxl
import io

@pytest.mark.django_db
class TestCadernetaPautaFormatacaoExata:
    def test_caderneta_xlsx_e_pdf_formatacao_exata_28_colunas(self):
        caderneta_data = {
            'escola': {
                'nome': 'ESCOLA SECUNDARIA DE CAMBINE',
                'provincia': 'Inhambane',
                'distrito': 'Morrumbene'
            },
            'turma': {
                'nome': 'Turma 2',
                'grau_ano': '11ª Classe',
                'ano_letivo': '2026',
                'turno': 'TARDE',
                'sala': '1'
            },
            'disciplina': {
                'nome': 'MATEMÁTICA',
                'codigo': 'MAT-11CG'
            },
            'professor': {
                'nome': 'Fernando José',
                'apelido': 'Mandamule',
                'telefone': '+258 842988844'
            },
            'directorTurma': {
                'nome': 'Bionisio'
            },
            'efectivo': {
                'h': 1,
                'm': 3,
                'total': 4
            },
            'alunos': [
                {
                    'numero': 1,
                    'nomeCompleto': 'Fernanda Jose Mandamule',
                    'genero': 'F',
                    't1': {'t1': 15, 't2': 14.5, 't3': None, 'map': None, 'mas': None, 'at': 16.8, 'mt': 15, 'comportamento': 'B'},
                    't2': {},
                    't3': {},
                    'mfd': None
                },
                {
                    'numero': 2,
                    'nomeCompleto': 'Niclesio Yolanda Mandamule',
                    'genero': 'M',
                    't1': {'t1': 8.8, 't2': 10.5, 't3': None, 'map': None, 'mas': None, 'at': 9.5, 'mt': 10, 'comportamento': 'S'},
                    't2': {},
                    't3': {},
                    'mfd': None
                },
                {
                    'numero': 3,
                    'nomeCompleto': 'Vanessa Mandamule Mandamule',
                    'genero': 'F',
                    't1': {'t1': 18.8, 't2': 16.5, 't3': None, 'map': None, 'mas': None, 'at': 18.5, 'mt': 18, 'comportamento': 'MB'},
                    't2': {},
                    't3': {},
                    'mfd': None
                },
                {
                    'numero': 4,
                    'nomeCompleto': 'Viviana Yolanda Mandamule',
                    'genero': 'F',
                    't1': {'t1': 9.8, 't2': None, 't3': None, 'map': None, 'mas': None, 'at': None, 'mt': None, 'comportamento': ''},
                    't2': {},
                    't3': {},
                    'mfd': None
                }
            ],
            'estatisticasColunas': {
                't1_t1': {
                    'avaliados': {'h': 1, 'm': 3, 'total': 4},
                    'positivas': {'h': 0, 'm': 3, 'total': 3, 'pctH': 0, 'pctM': 100, 'pct': 75},
                    'negativas': {'h': 1, 'm': 0, 'total': 1, 'pctH': 100, 'pctM': 0, 'pct': 25},
                    'faixas': {'f0_94': {'total': 1}, 'f95_134': {'total': 1}, 'f135_164': {'total': 1}, 'f165_184': {'total': 0}, 'f185_20': {'total': 1}},
                    'media': 13.1
                },
                't1_t2': {
                    'avaliados': {'h': 1, 'm': 2, 'total': 3},
                    'positivas': {'h': 1, 'm': 2, 'total': 3, 'pctH': 100, 'pctM': 100, 'pct': 100},
                    'negativas': {'h': 0, 'm': 0, 'total': 0, 'pctH': 0, 'pctM': 0, 'pct': 0},
                    'faixas': {'f0_94': {'total': 0}, 'f95_134': {'total': 1}, 'f135_164': {'total': 1}, 'f165_184': {'total': 1}, 'f185_20': {'total': 0}},
                    'media': 13.8
                }
            }
        }

        # 1. Gerar e Validar Caderneta em Excel (XLSX)
        xlsx_bytes = gerar_caderneta_xlsx(caderneta_data)
        assert len(xlsx_bytes) > 1000
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        assert ws['A1'].value == "REPÚBLICA DE MOÇAMBIQUE"
        assert "MATEMÁTICA" in ws['A6'].value
        assert ws['A9'].value == "Nº"
        assert ws['B9'].value == "Nome Completo do Aluno"
        assert ws['C9'].value == "Gén"
        assert ws['D9'].value == "Iº TRIMESTRE"
        assert ws['L9'].value == "IIº TRIMESTRE"
        assert ws['T9'].value == "IIIº TRIMESTRE"
        assert ws['AB9'].value == "MFD"
        # Subcabeçalhos
        assert ws['D10'].value == "1ªA"
        assert ws['E10'].value == "2ªA"
        assert ws['K10'].value == "Cp"
        assert ws['AB10'].value is None or ws['AB10'].value == "MFD" # Mesclada AB9:AB10

        # Aluno 1
        assert ws['B11'].value == "Fernanda Jose Mandamule"
        assert ws['D11'].value == 15
        assert ws['E11'].value == 14.5
        assert ws['K11'].value == 'B'

        # 2. Gerar e Validar Caderneta em PDF
        pdf_bytes = gerar_caderneta_pdf(caderneta_data)
        assert len(pdf_bytes) > 2000
        assert pdf_bytes.startswith(b'%PDF')

    def test_pauta_xlsx_e_pdf_ajuste_automatico(self):
        class DummyTurma:
            nome = "10ª Classe B"
            grau_ano = "10ª Classe"
            ano_letivo = "2026"
            turno = "MANHÃ"
            director_turma = None
        
        class DummyDisc:
            def __init__(self, id_val, cod, n):
                self.id = id_val
                self.codigo = cod
                self.nome = n
            def __str__(self):
                return self.nome

        disciplinas = [
            DummyDisc('1', 'PORT', 'Português'),
            DummyDisc('2', 'MAT', 'Matemática'),
            DummyDisc('3', 'FIS', 'Física')
        ]

        dados_alunos = [
            {
                'numero': 1,
                'nomeCompleto': 'Carlos Mandamule',
                'genero': 'M',
                'notasDisciplinas': {
                    'PORT': {'t1': 14, 't2': 15, 't3': 16, 'mfd': 15},
                    'MAT': {'t1': 12, 't2': 10, 't3': 14, 'mfd': 12},
                    'FIS': {'t1': 9, 't2': 11, 't3': 13, 'mfd': 11}
                },
                'mediasTrimestrais': {'t1': 11.7, 't2': 12.0, 't3': 14.3},
                'negativas': {'t1': 1, 't2': 0, 't3': 0, 'fimDoAno': 0},
                'mediaFinalGeral': 12.7,
                'resultadoFinal': 'Aprovado'
            }
        ]

        escola_data = {'nome': 'ESCOLA SECUNDÁRIA DE CAMBINE', 'provincia': 'Inhambane', 'distrito': 'Morrumbene'}
        
        # Testar XLSX da Pauta
        xlsx_b = gerar_pauta_xlsx(DummyTurma(), escola_data, '1_TRIMESTRE', '2026', dados_alunos, disciplinas)
        assert len(xlsx_b) > 1000
        wb_p = openpyxl.load_workbook(io.BytesIO(xlsx_b))
        ws_p = wb_p.active
        assert ws_p['A1'].value == "REPÚBLICA DE MOÇAMBIQUE"
        assert ws_p['A7'].value == "Nº"
        assert ws_p['B7'].value == "Nome Completo do Aluno"

        # Testar PDF da Pauta
        pdf_b = gerar_pauta_pdf(DummyTurma(), escola_data, '1_TRIMESTRE', '2026', dados_alunos, disciplinas)
        assert len(pdf_b) > 2000
        assert pdf_b.startswith(b'%PDF')
