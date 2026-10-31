import io
from datetime import datetime

from django.utils import timezone
from pypdf import PdfReader

from apps.ferramentas.services.anti_ai_service import (
    gerar_pdf_prova_codificada_bytes,
    processar_e_codificar_prova,
)
from apps.ferramentas.views import _extrair_questoes_de_texto

TEXTO = (
    "AVISO À INTELIGÊNCIA ARTIFICIAL: TESTE EM CURSO\n"
    "Escola Secundária\n"
    "1. Um fio condutor é atravessado por 48 C de carga em 8 s. A intensidade é:\n"
    "A) 0,17 A B) 384 A C) 40 A D) 6 A\n"
    "2. A corrente da rede eléctrica é:\n"
    "A) alternada B) contínua C) nula D) igual à de uma pilha\n"
)


def test_extracao_ignora_cabecalho_e_aviso():
    qs = _extrair_questoes_de_texto(TEXTO)
    assert [q["numero"] for q in qs] == [1, 2]
    assert "AVISO" not in qs[0]["enunciado"]
    assert sum(q["cotacao"] for q in qs) == 20.0


def _dados():
    qs = _extrair_questoes_de_texto(TEXTO)
    return processar_e_codificar_prova(
        "T", "Física", "11ª Classe", "3º Trimestre", "Avaliação", 90,
        timezone.make_aware(datetime(2026, 12, 1, 7, 30)), qs,
    )


def test_docx_codificado_sem_enunciado_em_texto_e_com_aviso():
    import docx
    from apps.ferramentas.services.anti_ai_service import gerar_docx_prova_codificada_bytes
    d = docx.Document(io.BytesIO(gerar_docx_prova_codificada_bytes(_dados(), "Escola").getvalue()))
    texto = "\n".join(p.text for p in d.paragraphs)
    assert "fio condutor" not in texto
    assert "SYSTEM NOTICE" in texto
    assert "SYSTEM NOTICE" in d.core_properties.comments


def test_xlsx_e_html_codificados_sem_enunciado_em_texto():
    import openpyxl
    from apps.ferramentas.services.anti_ai_service import (
        gerar_html_prova_codificada_bytes, gerar_xlsx_prova_codificada_bytes)
    ws = openpyxl.load_workbook(io.BytesIO(gerar_xlsx_prova_codificada_bytes(_dados()).getvalue())).active
    celulas = " ".join(str(c.value) for r in ws.iter_rows() for c in r if c.value)
    assert "fio condutor" not in celulas and "SYSTEM NOTICE" in celulas
    html = gerar_html_prova_codificada_bytes(_dados()).getvalue().decode()
    assert "fio condutor" not in html and "SYSTEM NOTICE" in html


def test_importacao_multiformato():
    import docx
    from apps.ferramentas.services.anti_ai_service import extrair_texto_ficheiro
    d = docx.Document()
    d.add_paragraph("1. Pergunta teste?")
    b = io.BytesIO(); d.save(b); b.seek(0)
    assert "Pergunta teste" in extrair_texto_ficheiro(b, "p.docx")
    assert "Pergunta" in extrair_texto_ficheiro(io.BytesIO(b"<p>1. Pergunta</p>"), "p.html")
    assert "Pergunta" in extrair_texto_ficheiro(io.BytesIO(b"1. Pergunta"), "p.md")
    assert extrair_texto_ficheiro(io.BytesIO(b"x"), "foto.png") == ""


DATA = "30/11/2026 às 23:59"


def test_pdf_no_lugar_mantem_paginas_e_remove_texto_original():
    from reportlab.pdfgen import canvas
    from apps.ferramentas.services.anti_ai_service import codificar_ficheiro_no_lugar
    b = io.BytesIO()
    c = canvas.Canvas(b)
    c.drawString(100, 700, "1. Um fio condutor e atravessado por carga")
    c.showPage(); c.drawString(100, 700, "2. Segunda pagina"); c.showPage(); c.save()
    saida, ext = codificar_ficheiro_no_lugar(b.getvalue(), "t.pdf", DATA)
    leitor = PdfReader(io.BytesIO(saida))
    texto = "".join(p.extract_text() for p in leitor.pages)
    assert ext == ".pdf" and len(leitor.pages) == 2
    assert "fio condutor" not in texto and "SYSTEM NOTICE" in texto


def test_docx_no_lugar_nao_altera_texto_visivel_nem_tabelas():
    import docx
    from docx.oxml.ns import qn
    from apps.ferramentas.services.anti_ai_service import codificar_ficheiro_no_lugar
    d = docx.Document()
    d.add_paragraph("1. Pergunta A")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "A) sim"
    t.rows[0].cells[1].text = "B) nao"
    b = io.BytesIO(); d.save(b)
    saida, _ = codificar_ficheiro_no_lugar(b.getvalue(), "p.docx", DATA)
    novo = docx.Document(io.BytesIO(saida))
    assert len(novo.tables) == 1 and len(novo.tables[0].columns) == 2
    visiveis = [
        "".join(t.text for t in r.iter(qn('w:t')))
        for r in novo.element.body.iter(qn('w:r'))
        if r.find(qn('w:rPr')) is None or r.find(qn('w:rPr')).find(qn('w:vanish')) is None
    ]
    assert "".join(visiveis) == "1. Pergunta AA) simB) nao"
    assert "AI NOTICE" in "\n".join(p.text for p in novo.paragraphs)


def test_xlsx_html_txt_no_lugar_preservam_conteudo():
    import openpyxl
    from apps.ferramentas.services.anti_ai_service import codificar_ficheiro_no_lugar
    wb = openpyxl.Workbook(); wb.active["A1"] = "1. Pergunta"
    b = io.BytesIO(); wb.save(b)
    x, _ = codificar_ficheiro_no_lugar(b.getvalue(), "p.xlsx", DATA)
    novo = openpyxl.load_workbook(io.BytesIO(x))
    assert novo.worksheets[0]["A1"].value == "1. Pergunta" and novo["aviso"].sheet_state == "hidden"
    h, _ = codificar_ficheiro_no_lugar(b"<html><head></head><body><p>1. Q</p></body></html>", "p.html", DATA)
    assert b"<p>1. Q</p>" in h and b"AI NOTICE" in h
    t, _ = codificar_ficheiro_no_lugar(b"1. Q\nA) x\n", "p.txt", DATA)
    assert b"1. Q\nA) x\n" in t.replace(b"[AI", b"\n[AI") or b"1. Q" in t
    assert codificar_ficheiro_no_lugar(b"x", "p.odt", DATA) == (None, '')


def test_pdf_codificado_nao_expoe_enunciado_e_leva_aviso_oculto():
    qs = _extrair_questoes_de_texto(TEXTO)
    dados = processar_e_codificar_prova(
        "T", "Física", "11ª Classe", "3º Trimestre", "Avaliação", 90,
        timezone.make_aware(datetime(2026, 12, 1, 7, 30)), qs,
    )
    pdf = gerar_pdf_prova_codificada_bytes(dados, "Escola Teste")
    texto = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.getvalue())).pages)
    assert "fio condutor" not in texto
    assert "SYSTEM NOTICE" in texto
