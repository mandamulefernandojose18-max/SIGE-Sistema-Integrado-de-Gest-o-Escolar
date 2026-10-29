import pytest
from rest_framework.test import APIClient
from django.utils import timezone
from apps.accounts.models import Usuario
from apps.tenants.models import Escola, Plano
from apps.turmas.models import Turma
from apps.alunos.models import Aluno
from apps.pagamentos.models import Pagamento
from apps.documentos.services.pdf_service import gerar_recibo_pdf
from apps.documentos.services.docx_service import gerar_recibo_docx
import io
import docx


@pytest.mark.django_db
class TestRecibosPagamentoStatus:
    def setup_method(self):
        self.client = APIClient()
        self.saas_plano = Plano.objects.create(nome='SAAS_PLANO_TESTE_RECIBO', preco=5000, duracao_dias=30)
        self.escola = Escola.objects.create(
            nome='Escola Secundária de Massinga',
            nif_cnpj='NUIT-99887766',
            provincia='Inhambane',
            distrito='Massinga',
            telefone='+258 84 123 4567',
            email='secretaria@massinga.edu.mz',
            status='ATIVA',
            plano=self.saas_plano
        )
        self.turma = Turma.objects.create(
            escola=self.escola,
            nome='10ª A',
            grau_ano='10ª',
            turno='DIURNO',
            sala='Sala 04'
        )
        self.aluno = Aluno.objects.create(
            escola=self.escola,
            turma=self.turma,
            matricula='MAT-2026-001',
            nome='Fernando',
            apelido='Mandamule',
            data_nascimento='2009-05-15',
            genero='M'
        )

        # 1. Pagamento NÃO PAGO (Pendente)
        self.pagamento_nao_pago = Pagamento.objects.create(
            escola=self.escola,
            aluno=self.aluno,
            descricao='Propina do Mês de Março 2026',
            mes_referencia='2026-03',
            valor=1500,
            valor_pago=0,
            status='PENDENTE',
            data_vencimento=timezone.now() + timezone.timedelta(days=10)
        )

        # 2. Pagamento PAGO (Liquidado)
        self.pagamento_pago = Pagamento.objects.create(
            escola=self.escola,
            aluno=self.aluno,
            descricao='Propina do Mês de Fevereiro 2026',
            mes_referencia='2026-02',
            valor=1500,
            valor_pago=1500,
            status='PAGO',
            metodo_pagamento='MPESA',
            data_vencimento=timezone.now() - timezone.timedelta(days=30),
            data_pagamento=timezone.now(),
            recibo_numero='REC-2026-00001'
        )

    def test_recibo_nao_pago_pdf(self):
        import re
        pdf_bytes = gerar_recibo_pdf(self.pagamento_nao_pago, self.escola)
        assert pdf_bytes is not None
        assert pdf_bytes.startswith(b'%PDF')
        num_paginas = len(re.findall(rb'/Type\s*/Page\b', pdf_bytes))
        assert num_paginas == 1, f"Recibo não pago deve ter exatamente 1 página (obtido: {num_paginas})"
        assert len(pdf_bytes) > 2000

    def test_recibo_pago_pdf(self):
        import re
        pdf_bytes = gerar_recibo_pdf(self.pagamento_pago, self.escola)
        assert pdf_bytes is not None
        assert pdf_bytes.startswith(b'%PDF')
        num_paginas = len(re.findall(rb'/Type\s*/Page\b', pdf_bytes))
        assert num_paginas == 1, f"Recibo pago deve ter exatamente 1 página (obtido: {num_paginas})"
        assert len(pdf_bytes) > 2000

    def test_recibo_nao_pago_docx(self):
        docx_bytes = gerar_recibo_docx(self.pagamento_nao_pago, self.escola)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = "\n".join([p.text for p in doc.paragraphs])
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    full_text += " " + cell.text

        assert "NÃO PAGO" in full_text
        assert "Aguardando Liquidação" in full_text
        assert "MASSINGA" in full_text.upper()
        assert "FERNANDO MANDAMULE" in full_text.upper()
        assert "Mil e quinhentos meticais" in full_text

    def test_recibo_pago_docx(self):
        docx_bytes = gerar_recibo_docx(self.pagamento_pago, self.escola)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = "\n".join([p.text for p in doc.paragraphs])
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    full_text += " " + cell.text

        assert "PAGO" in full_text
        assert "1500 MT" in full_text
        assert "Mil e quinhentos meticais" in full_text
        assert "MPESA" in full_text

    def test_recibo_impressao_view_nao_pago(self):
        resp = self.client.get(f'/api/v1/impressao/recibo/{self.pagamento_nao_pago.id}/')
        assert resp.status_code == 200
        data = resp.json()['data']
        assert data['is_pago'] is False
        assert data['status'] == 'NÃO PAGO'
        assert data['status_badge'] == 'danger'
        assert data['valor_extenso'] == 'Mil e quinhentos meticais'
        assert data['pagamento']['valor_pago'] == '0.00'
        assert data['escola']['nome'] == 'Escola Secundária de Massinga'
        assert data['aluno']['matricula'] == self.aluno.matricula

    def test_recibo_impressao_view_pago(self):
        resp = self.client.get(f'/api/v1/impressao/recibo/{self.pagamento_pago.id}/')
        assert resp.status_code == 200
        data = resp.json()['data']
        assert data['is_pago'] is True
        assert data['status'] == 'PAGO'
        assert data['status_badge'] == 'success'
        assert float(data['pagamento']['valor_pago']) == 1500.0
        assert data['valor_extenso'] == 'Mil e quinhentos meticais'
        assert 'Mil e quinhentos meticais' in data['pagamento']['valor_com_extenso']
        assert data['pagamento']['metodo_pagamento'] == 'MPESA'
        assert data['pagamento']['metodo_pagamento'] == 'MPESA'

    def test_recibo_pdf_endpoint(self):
        resp_nao_pago = self.client.get(f'/api/v1/impressao/recibo/{self.pagamento_nao_pago.id}/pdf')
        assert resp_nao_pago.status_code == 200
        assert resp_nao_pago['Content-Type'] == 'application/pdf'
        assert resp_nao_pago.content.startswith(b'%PDF')

        resp_pago = self.client.get(f'/api/v1/impressao/recibo/{self.pagamento_pago.id}/pdf')
        assert resp_pago.status_code == 200
        assert resp_pago['Content-Type'] == 'application/pdf'
        assert resp_pago.content.startswith(b'%PDF')

    def test_aluno_me_pagamentos_e_recibo_pdf(self):
        """Valida que o aluno acede aos seus pagamentos por período e descarrega o recibo em PDF"""
        user_aluno = Usuario.objects.create_user(
            email='aluno.recibo@massinga.edu.mz',
            password='SenhaSegura123!',
            nome='Fernando Mandamule',
            role='ALUNO',
            aluno_id=self.aluno.id,
            escola=self.escola,
            codigo_acesso='ALU-2026-9988'
        )
        self.client.force_authenticate(user=user_aluno)

        # 1. Consulta de pagamentos do aluno (sem filtro de ano)
        res_pagamentos = self.client.get('/api/v1/alunos/me/pagamentos')
        assert res_pagamentos.status_code == 200
        dados = res_pagamentos.json()
        assert dados['success'] is True
        assert len(dados['data']) == 2

        # 2. Consulta filtrando pelo período/ano letivo
        res_ano = self.client.get('/api/v1/alunos/me/pagamentos?ano=2026')
        assert res_ano.status_code == 200
        assert res_ano.json()['success'] is True
        assert len(res_ano.json()['data']) == 2

        # 3. Identifica o pagamento efetuado e descarrega o recibo exclusivamente em PDF
        pag_pago = next(p for p in dados['data'] if p['status'] == 'PAGO')
        assert pag_pago['id'] == str(self.pagamento_pago.id)

        res_pdf = self.client.get(f"/api/v1/impressao/recibo/{pag_pago['id']}/pdf")
        assert res_pdf.status_code == 200
        assert res_pdf['Content-Type'] == 'application/pdf'
        assert res_pdf.content.startswith(b'%PDF')
        assert 'filename="Recibo' in res_pdf.headers.get('Content-Disposition', '')

