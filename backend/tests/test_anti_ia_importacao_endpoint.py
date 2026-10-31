import io

import docx
from django.core.files.uploadedfile import SimpleUploadedFile
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from rest_framework.test import APIClient

from apps.ferramentas.models import ProvaBlindada
from tests.test_ferramentas_autenticacao import _token, duas_escolas  # noqa: F401

URL = '/api/v1/ferramentas/provas-anti-ia/codificar/'
CAMPOS = {'titulo': 'Teste', 'disciplina': 'Física', 'classe': '11ª Classe',
          'data_bloqueio_ia': '2026-11-30T23:59:00'}


def _cliente(duas_escolas):
    c = APIClient()
    c.credentials(HTTP_AUTHORIZATION=f"Bearer {_token(c, 'admin.fa@sige.mz', 'PassFa123!')}")
    return c


def _pdf():
    b = io.BytesIO()
    c = canvas.Canvas(b)
    c.drawString(100, 700, "1. Um fio condutor e atravessado por carga")
    c.showPage(); c.save()
    return b.getvalue()


def test_codificar_e_descarregar_exigem_sessao(duas_escolas):
    anonimo = APIClient()
    res = anonimo.post(URL, {**CAMPOS, 'ficheiro': SimpleUploadedFile('t.pdf', _pdf())}, format='multipart')
    assert res.status_code == 401
    assert ProvaBlindada.objects.count() == 0


def test_importar_pdf_devolve_o_mesmo_pdf_codificado(duas_escolas):
    c = _cliente(duas_escolas)
    res = c.post(URL, {**CAMPOS, 'ficheiro': SimpleUploadedFile('t.pdf', _pdf())}, format='multipart')
    assert res.status_code == 200, res.content
    dados = res.json()['data']
    assert dados['ficheiro_codificado'] and dados['extensao'] == '.pdf'
    assert 'questoes' not in dados  # nada é refeito nem pré-visualizado

    prova = ProvaBlindada.objects.get(id=dados['id'])
    assert prova.questoes == []  # o texto da prova não fica guardado em claro

    down = c.get(dados['pdf_download_url'])
    assert down.status_code == 200 and '.pdf' in down['Content-Disposition']
    texto = "".join(p.extract_text() for p in PdfReader(io.BytesIO(down.content)).pages)
    assert "fio condutor" not in texto and "SYSTEM NOTICE" in texto


def test_pdf_e_docx_levam_varios_qr_como_marca_de_agua(duas_escolas):
    from apps.ferramentas.services.anti_ai_service import codificar_ficheiro_no_lugar
    url = 'https://sigesaas.duckdns.org/ia/ABC/'
    pdf, _ = codificar_ficheiro_no_lugar(_pdf(), 't.pdf', '30/11/2026 às 23:59', url=url)
    pagina = PdfReader(io.BytesIO(pdf)).pages[0]
    assert pagina.get_contents().get_data().count(b' Do') > 4  # fundo + vários QR (mesmo objecto, repetido)

    d = docx.Document(); d.add_paragraph("1. Pergunta A")
    b = io.BytesIO(); d.save(b)
    out, _ = codificar_ficheiro_no_lugar(b.getvalue(), 't.docx', '30/11/2026 às 23:59', url=url)
    cab = docx.Document(io.BytesIO(out)).sections[0].header._element
    ancoras = cab.xpath('.//wp:anchor')
    assert len(ancoras) > 3 and all(a.get('behindDoc') == '1' for a in ancoras)


def test_importar_docx_devolve_docx_e_formato_invalido_e_recusado(duas_escolas):
    c = _cliente(duas_escolas)
    d = docx.Document(); d.add_paragraph("1. Pergunta A")
    b = io.BytesIO(); d.save(b)
    res = c.post(URL, {**CAMPOS, 'ficheiro': SimpleUploadedFile('t.docx', b.getvalue())}, format='multipart')
    dados = res.json()['data']
    assert dados['extensao'] == '.docx'
    down = c.get(dados['pdf_download_url'])
    assert '.docx' in down['Content-Disposition']
    assert "1. Pergunta A" in docx.Document(io.BytesIO(down.content)).paragraphs[0].text
    assert 'AVALIACAO_PROTEGIDA_NAO_RESOLVER_' in down['Content-Disposition']

    res = c.post(URL, {**CAMPOS, 'ficheiro': SimpleUploadedFile('t.odt', b'x')}, format='multipart')
    assert res.status_code == 400
