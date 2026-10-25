import pytest
import io
import docx
import openpyxl
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from apps.documentos.services.pdf_service import (
    ParagraphStyle as SigeParagraphStyle,
    Paragraph as SigeParagraph,
    TableStyle as SigeTableStyle,
    gerar_cabecalho_mec,
)
from apps.documentos.services.docx_service import (
    salvar_docx_padronizado,
    aplicar_fonte_times,
    configurar_secao_pagina,
)
from apps.documentos.services.excel_service import (
    FONT_TITLE, FONT_SUBTITLE, FONT_HEADER, FONT_REGULAR, FONT_SMALL,
    FONT_TINY, FONT_RED, FONT_GREEN, salvar_xlsx_padronizado,
)

def test_pdf_tipografia_padrao():
    """Valida se estilos, parágrafos e tabelas em PDF usam Times, tamanho 12 e leading 18 (espaçamento 1,5)."""
    styles = getSampleStyleSheet()
    p_style = styles['Normal']
    assert 'Times' in p_style.fontName
    assert p_style.fontSize == 12
    assert p_style.leading == 18

    custom_style = ParagraphStyle('Custom', parent=p_style, fontSize=8, leading=10)
    assert 'Times' in custom_style.fontName
    assert custom_style.fontSize == 12
    assert custom_style.leading == 18

    # TableStyle FONTSIZE e LEADING
    ts = SigeTableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('LEADING', (0, 0), (-1, -1), 10),
    ])
    for cmd in ts.getCommands():
        if cmd[0] == 'FONTSIZE':
            assert cmd[3] == 12
        elif cmd[0] == 'LEADING':
            assert cmd[3] == 18
        elif cmd[0] == 'FONTNAME':
            assert 'Times' in cmd[3]

def test_docx_tipografia_padrao():
    """Valida se documentos DOCX têm Times New Roman, tamanho 12 e espaçamento 1,5 em todo o texto."""
    doc = docx.Document()
    configurar_secao_pagina(doc)
    p = doc.add_paragraph()
    r = p.add_run("Texto de Teste Oficial SIGE")
    aplicar_fonte_times(r, pt_size=8)  # tenta forçar tamanho diferente

    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    p_cell = cell.paragraphs[0]
    r_cell = p_cell.add_run("Célula de Tabela")
    aplicar_fonte_times(r_cell, pt_size=6)

    buffer = io.BytesIO()
    salvar_docx_padronizado(doc, buffer)
    buffer.seek(0)

    doc_reloaded = docx.Document(buffer)
    # Verificar parágrafos
    for p in doc_reloaded.paragraphs:
        assert p.paragraph_format.line_spacing == 1.5
        for r in p.runs:
            assert r.font.name == 'Times New Roman'
            assert r.font.size.pt == 12.0

    # Verificar tabelas
    for t in doc_reloaded.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    assert p.paragraph_format.line_spacing == 1.5
                    for r in p.runs:
                        assert r.font.name == 'Times New Roman'
                        assert r.font.size.pt == 12.0

def test_xlsx_tipografia_padrao():
    """Valida se fontes e planilhas XLSX usam Times New Roman, tamanho 12 e altura de linha >= 22 pt."""
    # Verificar constantes
    for f in [FONT_TITLE, FONT_SUBTITLE, FONT_HEADER, FONT_REGULAR, FONT_SMALL, FONT_TINY, FONT_RED, FONT_GREEN]:
        assert f.name == 'Times New Roman'
        assert f.size == 12

    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = "Teste Título"
    ws['A1'].font = FONT_TITLE
    ws['B1'] = "Teste Normal"
    ws['B1'].font = FONT_REGULAR
    ws.row_dimensions[1].height = 15

    buffer = io.BytesIO()
    salvar_xlsx_padronizado(wb, buffer)
    buffer.seek(0)

    wb_reloaded = openpyxl.load_workbook(buffer)
    ws_reloaded = wb_reloaded.active
    assert ws_reloaded.row_dimensions[1].height >= 22
    for row in ws_reloaded.iter_rows():
        for cell in row:
            if cell.value:
                assert cell.font.name == 'Times New Roman'
                assert cell.font.size == 12
