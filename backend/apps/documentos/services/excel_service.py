import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def gerar_pauta_xlsx(turma, escola, periodo, ano_letivo, dados_alunos, disciplinas) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Pauta {turma.nome}"

    # Estilos
    font_header = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
    font_bold = Font(name='Arial', size=10, bold=True)
    font_normal = Font(name='Arial', size=9)
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='BDC3C7'),
        right=Side(style='thin', color='BDC3C7'),
        top=Side(style='thin', color='BDC3C7'),
        bottom=Side(style='thin', color='BDC3C7')
    )

    # Cabeçalho da Instituição
    ws.merge_cells('A1:J1')
    ws['A1'] = "REPÚBLICA DE MOÇAMBIQUE — MINISTÉRIO DA EDUCAÇÃO E DESENVOLVIMENTO HUMANO"
    ws['A1'].font = font_bold
    ws['A1'].alignment = align_center

    ws.merge_cells('A2:J2')
    ws['A2'] = f"{escola.nome} | PAUTA OFICIAL DE APROVEITAMENTO — {periodo.replace('_', ' ')} / {ano_letivo}"
    ws['A2'].font = font_bold
    ws['A2'].alignment = align_center

    ws.merge_cells('A3:J3')
    ws['A3'] = f"Turma: {turma.nome} | Grau: {turma.grau_ano} | Turno: {turma.turno}"
    ws['A3'].font = font_normal
    ws['A3'].alignment = align_center

    # Linha da Tabela
    headers = ['Nº', 'Nome Completo'] + [d.codigo for d in disciplinas] + ['Média Geral', 'Negativas', 'Resultado']
    ws.append([])
    ws.append(headers)
    header_row_idx = 5

    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    # Dados dos Alunos
    for idx, aluno in enumerate(dados_alunos, start=1):
        row = [idx, aluno['nome']]
        for d in disciplinas:
            nota = (aluno['notas'].get(str(d.id)) or aluno['notas'].get(d.id) or {}).get('media_final', '-')
            row.append(nota)
        row.append(aluno.get('media_geral', '-'))
        row.append(aluno.get('total_negativas', 0))
        row.append(aluno.get('resultado', 'Aprovado'))
        ws.append(row)

        curr_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            cell.alignment = align_left if col_idx == 2 else align_center

    # Ajuste de largura das colunas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
